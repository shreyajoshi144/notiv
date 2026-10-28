"""Foundation tests: storage, upload validation, assets, authorization,
evidence, durable jobs (state transitions, retry, stage isolation, recovery)."""
import os
import shutil
import subprocess
import tempfile
import time
import unittest
import zipfile

from core.lecture import jobs, pipeline, validation
from core.lecture.align import Segment
from core.lecture.errors import NotFoundError, PermanentError, ValidationError
from core.lecture.evidence import EvidenceService
from core.lecture.make_test_lecture import make_lecture
from core.lecture.storage import LocalStorage
from core.lecture.testkit import Env, segments
from core.lecture.transcribe import TranscriptionError

SEGS = [Segment(1.0, 4.0, "Gradient descent uses the gradient."), Segment(7.0, 12.0, "The learning rate is the step size.")]


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.video, cls.pdf, _ = make_lecture(os.path.join(cls.tmp, "fx"))
        cls.env = Env(cls.tmp)
        cls.repo = cls.env.repo

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def write(self, name, data: bytes) -> str:
        p = os.path.join(self.tmp, name)
        with open(p, "wb") as f:
            f.write(data)
        return p


class TestStorage(Base):
    def setUp(self):
        self.st = LocalStorage(os.path.join(self.tmp, "st_" + self._testMethodName), signing_secret="s3cret")

    def test_put_get_exists_delete(self):
        src = self.write("s.bin", b"hello")
        self.assertEqual(self.st.put("a/b/c.bin", src), 5)
        self.assertTrue(self.st.exists("a/b/c.bin"))
        self.assertEqual(self.st.read_bytes("a/b/c.bin"), b"hello")
        with self.st.local_path("a/b/c.bin") as p:
            self.assertEqual(open(p, "rb").read(), b"hello")
        self.st.delete("a/b/c.bin")
        self.assertFalse(self.st.exists("a/b/c.bin"))
        self.st.delete("a/b/c.bin")                     # idempotent

    def test_traversal_and_bad_keys_rejected(self):
        for bad in ("../x", "a/../../x", "/abs", "a//b", "a\\b", ""):
            with self.assertRaises(ValidationError, msg=bad):
                self.st.put_bytes(bad, b"x")

    def test_delete_prefix(self):
        for k in ("p/1", "p/2/3", "q/1"):
            self.st.put_bytes(k, b"x")
        self.assertEqual(self.st.delete_prefix("p"), 2)
        self.assertTrue(self.st.exists("q/1"))
        self.assertFalse(self.st.exists("p/1"))

    def test_signed_url_verify_expiry_and_tamper(self):
        url = self.st.signed_url("a/b.bin", expires_in=60)
        qs = dict(kv.split("=") for kv in url.split("?")[1].split("&"))
        exp, sig = int(qs["exp"]), qs["sig"]
        self.assertTrue(self.st.verify_signed("a/b.bin", exp, sig))
        self.assertFalse(self.st.verify_signed("a/other.bin", exp, sig))          # different key
        self.assertFalse(self.st.verify_signed("a/b.bin", exp + 1, sig))          # tampered expiry
        self.assertFalse(self.st.verify_signed("a/b.bin", exp, sig, now=exp + 1))  # expired


class TestValidation(Base):
    def test_valid_types_detected_from_content(self):
        v = validation.validate_file(self.video, "lecture.mp4", "video/mp4")
        self.assertEqual((v.type, v.mime_type), ("video", "video/mp4"))
        self.assertAlmostEqual(v.duration_seconds, 27, delta=1)
        self.assertEqual(v.checksum, validation.sha256_file(self.video))
        self.assertEqual(validation.validate_file(self.pdf, "d.pdf", "application/pdf").type, "pdf")
        audio = os.path.join(self.tmp, "only.mp4")                                  # audio inside an .mp4 name
        subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", self.video, "-vn", "-c:a", "aac", "-f", "ipod", audio], check=True, timeout=30)
        self.assertEqual(validation.validate_file(audio, "only.mp4").type, "audio")

    def test_pptx_and_image(self):
        from pptx import Presentation
        p = os.path.join(self.tmp, "ok.pptx"); pr = Presentation(); pr.slides.add_slide(pr.slide_layouts[0]); pr.save(p)
        self.assertEqual(validation.validate_file(p, "ok.pptx").type, "pptx")
        import struct, zlib     # Sprint 5: image validation now requires a readable IHDR (dimension cap), so use a real header
        ihdr = struct.pack(">IIBBBBB", 10, 10, 8, 2, 0, 0, 0)
        png = self.write("i.png", b"\x89PNG\r\n\x1a\n" + struct.pack(">I", 13) + b"IHDR" + ihdr + struct.pack(">I", zlib.crc32(b"IHDR" + ihdr)))
        with self.assertRaises(ValidationError):
            validation.validate_file(png, "i.png", "image/png")          # images are no longer a supported upload

    def test_content_must_match_extension(self):
        pdf_as_mp4 = shutil.copyfile(self.pdf, os.path.join(self.tmp, "fake.mp4"))
        mp4_as_pdf = shutil.copyfile(self.video, os.path.join(self.tmp, "fake.pdf"))
        text_as_pptx = self.write("fake.pptx", b"just text, not a zip")
        zip_not_pptx = os.path.join(self.tmp, "z.pptx")
        with zipfile.ZipFile(zip_not_pptx, "w") as z:
            z.writestr("readme.txt", "x")
        for path, name in ((pdf_as_mp4, "fake.mp4"), (mp4_as_pdf, "fake.pdf"), (text_as_pptx, "fake.pptx"),
                           (zip_not_pptx, "z.pptx"), (self.write("e.png", b"notpng" * 10), "e.png")):
            with self.assertRaises(ValidationError, msg=name):
                validation.validate_file(path, name)

    def test_mime_size_empty_unknown_ext(self):
        with self.assertRaises(ValidationError):
            validation.validate_file(self.pdf, "d.pdf", "video/mp4")                # wrong declared MIME
        with self.assertRaises(ValidationError):
            validation.validate_file(self.write("empty.pdf", b""), "empty.pdf")
        with self.assertRaises(ValidationError):
            validation.validate_file(self.write("x.exe", b"MZ"), "x.exe")
        old = validation.MAX_BYTES
        validation.MAX_BYTES = 10
        try:
            with self.assertRaises(ValidationError):
                validation.validate_file(self.pdf, "d.pdf")
        finally:
            validation.MAX_BYTES = old

    def test_duration_limit(self):
        old = validation.MAX_DURATION_SEC
        validation.MAX_DURATION_SEC = 5
        try:
            with self.assertRaises(ValidationError):
                validation.validate_file(self.video, "lecture.mp4")
        finally:
            validation.MAX_DURATION_SEC = old


class TestAssetsAndOwnership(Base):
    def test_asset_record_and_storage_key_ignores_filename(self):
        lid = self.env.lectures.create(1, "L")
        a = self.env.assets.add_from_file(1, lid, self.video, "../../etc/passwd/My Lecture.MP4", "video/mp4")
        self.assertEqual((a["type"], a["status"], a["lecture_id"], a["uploader_id"]), ("video", "ready", lid, 1))
        self.assertEqual(a["original_filename"], "My Lecture.MP4")                   # basename only, display use
        self.assertRegex(a["storage_key"], rf"^lectures/{lid}/assets/[0-9a-f]{{32}}\.mp4$")
        self.assertNotIn("passwd", a["storage_key"])
        self.assertTrue(a["checksum"] and a["size_bytes"] > 0 and a["duration_seconds"] > 0)
        self.assertTrue(self.env.storage.exists(a["storage_key"]))

    def test_rejected_upload_stores_nothing(self):
        lid = self.env.lectures.create(1, "L")
        bad = self.write("bad.mp4", b"not media at all")
        with self.assertRaises(ValidationError):
            self.env.assets.add_from_file(1, lid, bad, "bad.mp4")
        self.assertEqual(self.repo.list_assets(lid), [])
        self.assertFalse(os.path.exists(os.path.join(self.env.storage.root, "lectures", str(lid), "assets")))

    def test_other_user_cannot_add_or_read_assets(self):
        lid = self.env.lectures.create(1, "mine")
        a = self.env.assets.add_from_file(1, lid, self.pdf, "d.pdf")
        with self.assertRaises(NotFoundError):
            self.env.assets.add_from_file(2, lid, self.pdf, "d.pdf")
        with self.assertRaises(NotFoundError):
            self.env.assets.list(2, lid)
        with self.assertRaises(NotFoundError):
            self.env.assets.signed_url(2, a["id"])
        self.assertIn("sig=", self.env.assets.signed_url(1, a["id"]))

    def test_workspace_member_gets_access_non_member_does_not(self):
        ws = self.repo.create_workspace(1, "team")
        lid = self.env.lectures.create(1, "shared", workspace_id=ws)
        with self.assertRaises(NotFoundError):
            self.env.assets.list(3, lid)
        self.repo.add_member(ws, 3)
        self.assertEqual(self.env.assets.list(3, lid), [])

    def test_cannot_create_lecture_in_foreign_workspace_or_course(self):
        ws = self.repo.create_workspace(1, "private")
        with self.assertRaises(NotFoundError):
            self.env.lectures.create(2, "x", workspace_id=ws)
        c = self.env.lectures.create_course(1, "course", ws)
        with self.assertRaises(NotFoundError):
            self.env.lectures.create(2, "x", course_id=c)

    def test_delete_lecture_removes_everything(self):
        lid = self.env.lecture(1, self.video, self.pdf)
        self.env.process(lid, segments(*SEGS))
        self.assertTrue(os.path.isdir(os.path.join(self.env.storage.root, "lectures", str(lid))))
        with self.assertRaises(NotFoundError):
            self.env.assets.delete_lecture(2, lid)
        self.env.assets.delete_lecture(1, lid)
        self.assertFalse(os.path.exists(os.path.join(self.env.storage.root, "lectures", str(lid))))
        self.assertIsNone(self.repo.get_lecture(lid))
        self.assertEqual((self.repo.list_assets(lid), self.repo.list_evidence(lid), self.repo.list_jobs(lid)), ([], [], []))
        self.assertEqual(self.repo.search_evidence([lid], "gradient"), [])


class TestEvidenceAndAuthz(Base):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.lid = cls.env.lecture(1, cls.video, cls.pdf)
        cls.env.process(cls.lid, segments(*SEGS))
        cls.ev = EvidenceService(cls.repo)

    def test_evidence_points_to_segment_slide_asset_and_time(self):
        assets = {a["type"]: a for a in self.repo.list_assets(self.lid)}
        rows = {e["source_type"]: e for e in self.repo.list_evidence(self.lid)}
        t = next(e for e in self.repo.list_evidence(self.lid) if e["source_type"] == "transcript" and e["start"] == 7.0)
        seg = next(s for s in self.repo.viewer_data(self.lid)["segments"] if s["start"] == 7.0)
        self.assertEqual((t["source_id"], t["slide_index"], t["end"], t["source_asset_id"]),
                         (seg["id"], 2, 12.0, assets["video"]["id"]))
        slide_id = {s["idx"]: s["id"] for s in self.repo.viewer_data(self.lid)["slides"]}[2]
        self.assertEqual(t["slide_id"], slide_id)
        self.assertGreater(t["confidence"], 0.9)
        self.assertTrue(all(e["created_at"] for e in self.repo.list_evidence(self.lid)))

    def test_resolve_gives_asset_and_seek_target(self):
        eid = next(e["id"] for e in self.repo.list_evidence(self.lid) if e["start"] == 7.0)
        r = self.ev.resolve(1, eid)
        self.assertEqual((r["seek_seconds"], r["slide_index"], r["asset"]["type"]), (7.0, 2, "video"))
        self.assertTrue(r["asset"]["storage_key"].startswith(f"lectures/{self.lid}/assets/"))

    def test_every_access_path_is_owner_scoped(self):
        eid = self.repo.list_evidence(self.lid)[0]["id"]
        job = self.repo.list_jobs(self.lid)[0]["id"]
        from core.lecture import authz
        for fn, arg in ((self.ev.get, eid), (self.ev.resolve, eid), (self.ev.list, self.lid)):
            with self.assertRaises(NotFoundError):
                fn(2, arg)
        with self.assertRaises(NotFoundError):
            authz.get_job(self.repo, 2, job)
        with self.assertRaises(NotFoundError):
            authz.get_lecture(self.repo, 2, self.lid)
        with self.assertRaises(NotFoundError):
            authz.get_lecture(self.repo, 1, 99999)               # missing looks identical
        with self.assertRaises(NotFoundError):
            self.env.lectures.status(2, self.lid)
        self.assertEqual(authz.accessible_lecture_ids(self.repo, 2, [self.lid]), [])
        self.assertEqual(authz.accessible_lecture_ids(self.repo, 1, [self.lid]), [self.lid])

    def test_new_derived_evidence_is_indexed_and_scoped(self):
        eid = self.ev.create(1, self.lid, "note", 1, "midterm will cover eigenvalues", start=7.0, slide_index=2)
        self.assertEqual([h["evidence_id"] for h in self.repo.search(self.lid, "eigenvalues")], [eid])
        with self.assertRaises(NotFoundError):
            self.ev.create(2, self.lid, "note", 1, "intruder")


class FakeClock:
    def __init__(self):
        self.t = 1000.0

    def __call__(self):
        return self.t


class TestJobs(Base):
    WATCHDOG_SECONDS = 240

    def setUp(self):
        # A stuck test must fail with a traceback showing WHERE, not hang the run.
        import signal
        if hasattr(signal, "SIGALRM"):
            def _timeout(signum, frame):
                raise TimeoutError(f"{self.id()} exceeded {self.WATCHDOG_SECONDS}s (hung)")
            self._old_handler = signal.signal(signal.SIGALRM, _timeout)
            signal.alarm(self.WATCHDOG_SECONDS)

    def tearDown(self):
        import signal
        if hasattr(signal, "SIGALRM"):
            signal.alarm(0)
            signal.signal(signal.SIGALRM, self._old_handler)

    def _setup(self, transcriber, max_attempts=3, clock=None):
        lid = self.env.lecture(1, self.video, self.pdf)
        js = pipeline.make_job_service(self.repo, self.env.storage, transcriber, backoff_base=10,
                                       **({"clock": clock} if clock else {}))
        return lid, js, js.enqueue(lid, max_attempts)

    def stages(self, jid):
        return {s["stage"]: s for s in self.repo.list_stages(jid)}

    def test_happy_path_states_and_all_ten_stages(self):
        lid, js, jid = self._setup(segments(*SEGS))
        j = self.repo.get_job(jid)
        self.assertEqual((j["status"], j["attempt"], j["started_at"], j["type"]), ("queued", 0, None, "process_lecture"))
        self.assertEqual(self.repo.get_lecture(lid)["status"], "queued")
        self.assertEqual(js.run_next().status, jobs.COMPLETED)
        j = self.repo.get_job(jid)
        self.assertEqual((j["status"], j["attempt"], j["progress"], j["error"]), ("completed", 1, 100, None))
        self.assertTrue(j["started_at"] and j["finished_at"] and j["created_at"])
        self.assertEqual(list(self.stages(jid)), pipeline.STAGE_NAMES)
        self.assertTrue(all(s["status"] == "completed" for s in self.stages(jid).values()))
        self.assertEqual(self.repo.get_lecture(lid)["status"], "done")
        self.assertIsNone(js.run_next())                          # nothing left to run

    def test_skipped_stages_for_audio_only(self):
        audio = os.path.join(self.tmp, "j.mp3")
        subprocess.run(["ffmpeg", "-y", "-v", "error", "-i", self.video, "-vn", "-c:a", "libmp3lame", audio], check=True, timeout=30)
        lid = self.env.lecture(1, audio)
        js = self.env.jobs(segments(*SEGS))
        jid = js.enqueue(lid)
        js.run_next()
        st = self.stages(jid)
        self.assertEqual([n for n, s in st.items() if s["status"] == "skipped"],
                         ["PROCESS_SLIDES", "EXTRACT_FRAMES", "DETECT_SLIDE_CHANGES"])

    def test_transient_failure_retries_with_backoff_and_resumes_without_rerunning(self):
        calls = {"n": 0}

        def flaky(wav, lang, prog, dur):
            calls["n"] += 1
            if calls["n"] == 1:
                raise TranscriptionError("model download hiccup")
            return list(SEGS)
        clock = FakeClock()
        lid, js, jid = self._setup(flaky, clock=clock)
        r = js.run_next()
        self.assertEqual(r.status, jobs.RETRYING)
        j = self.repo.get_job(jid)
        self.assertEqual((j["status"], j["attempt"]), ("retrying", 1))
        self.assertIn("[TRANSCRIBE] TranscriptionError", j["error"])
        self.assertEqual(j["run_after"], clock.t + 10)             # backoff_base * 2**0
        st = self.stages(jid)
        self.assertEqual((st["EXTRACT_AUDIO"]["status"], st["TRANSCRIBE"]["status"]), ("completed", "failed"))
        self.assertIsNone(js.run_next())                           # not due yet
        clock.t += 11
        self.assertEqual(js.run_next().status, jobs.COMPLETED)
        st = self.stages(jid)
        self.assertEqual(st["EXTRACT_AUDIO"]["attempt"], 1)         # NOT re-run
        self.assertEqual(st["PROBE"]["attempt"], 1)
        self.assertEqual(st["TRANSCRIBE"]["attempt"], 2)            # only the failed stage re-ran
        self.assertEqual(self.repo.get_job(jid)["attempt"], 2)
        self.assertEqual(self.repo.get_lecture(lid)["status"], "done")
        self.assertIsNone(self.repo.get_lecture(lid)["error"])

    def test_late_stage_failure_never_retranscribes(self):
        calls = {"n": 0}

        def counting(wav, lang, prog, dur):
            calls["n"] += 1
            return list(SEGS)
        clock = FakeClock()
        lid, js, jid = self._setup(counting, clock=clock)
        real, fail = self.repo.rebuild_index, {"once": True}

        def flaky_index(lecture_id):
            if fail.pop("once", False):
                raise RuntimeError("index backend down")
            return real(lecture_id)
        self.repo.rebuild_index = flaky_index
        try:
            self.assertEqual(js.run_next().status, jobs.RETRYING)
            self.assertEqual(self.stages(jid)["INDEX"]["status"], "failed")
            self.assertEqual(self.stages(jid)["BUILD_EVIDENCE"]["status"], "completed")
            clock.t += 11
            self.assertEqual(js.run_next().status, jobs.COMPLETED)
        finally:
            del self.repo.rebuild_index
        self.assertEqual(calls["n"], 1)                             # transcribed exactly once
        self.assertEqual(self.stages(jid)["TRANSCRIBE"]["attempt"], 1)
        self.assertEqual(len(self.repo.search(lid, "gradient")), 1)

    def test_permanent_failure_does_not_retry(self):
        def bad(*a):
            raise PermanentError("unsupported language model")
        lid, js, jid = self._setup(bad)
        self.assertEqual(js.run_next().status, jobs.FAILED)
        j = self.repo.get_job(jid)
        self.assertEqual((j["status"], j["attempt"]), ("failed", 1))
        self.assertIsNone(js.run_next())
        self.assertEqual(self.repo.get_lecture(lid)["status"], "failed")

    def test_attempts_exhausted_then_manual_retry_resumes(self):
        calls = {"n": 0}

        def always_then_ok(wav, lang, prog, dur):
            calls["n"] += 1
            if calls["n"] <= 2:
                raise TranscriptionError("boom")
            return list(SEGS)
        clock = FakeClock()
        lid, js, jid = self._setup(always_then_ok, max_attempts=2, clock=clock)
        self.assertEqual(js.run_next().status, jobs.RETRYING)
        clock.t += 100
        self.assertEqual(js.run_next().status, jobs.FAILED)
        j = self.repo.get_job(jid)
        self.assertEqual((j["status"], j["attempt"], j["max_attempts"]), ("failed", 2, 2))
        self.assertEqual(self.repo.get_lecture(lid)["status"], "failed")
        self.assertTrue(js.retry(jid))
        self.assertFalse(js.retry(jid + 999))
        self.assertEqual(js.run_next().status, jobs.COMPLETED)
        self.assertEqual(self.stages(jid)["EXTRACT_AUDIO"]["attempt"], 1)
        self.assertEqual(self.repo.get_job(jid)["attempt"], 3)
        self.assertFalse(js.retry(jid))                             # completed jobs can't be retried

    def test_cancel_queued_and_cancel_between_stages(self):
        lid, js, jid = self._setup(segments(*SEGS))
        self.assertTrue(js.cancel(jid))
        self.assertEqual(self.repo.get_job(jid)["status"], "cancelled")
        self.assertIsNone(js.run_next())
        self.assertFalse(js.cancel(jid))

        lid2, js2, jid2 = self._setup(segments(*SEGS))
        orig = self.repo.set_stage

        def cancelling(job_id, stage, status, **kw):
            orig(job_id, stage, status, **kw)
            if stage == "EXTRACT_AUDIO" and status == "completed":
                js2.cancel(job_id)
        self.repo.set_stage = cancelling
        try:
            self.assertEqual(js2.run_next().status, jobs.CANCELLED)
        finally:
            del self.repo.set_stage
        self.assertNotIn("TRANSCRIBE", self.stages(jid2))           # stopped before the next stage
        self.assertEqual(self.repo.get_lecture(lid2)["status"], "failed")

    def test_crashed_worker_job_is_recovered_and_resumes(self):
        # Own database: recover_stale_jobs()/claim are table-wide operations, so this test must not
        # share a jobs table with other tests (no cross-test claims, no foreign stale jobs).
        env = Env(os.path.join(self.tmp, "crash_env"))
        repo = env.repo
        clock = FakeClock()
        lid = env.lecture(1, self.video, self.pdf)
        js = pipeline.make_job_service(repo, env.storage, segments(*SEGS), backoff_base=0, clock=clock)
        jid = js.enqueue(lid)
        claimed = repo.claim_next_job(clock.t, job_id=jid)          # worker claims, completes VALIDATE, then dies
        repo.set_stage(jid, "VALIDATE", "running")
        repo.set_stage(jid, "VALIDATE", "completed", output=jobs_validate_output(repo, lid))
        self.assertEqual((claimed["status"], claimed["attempt"]), ("running", 1))
        self.assertIsNone(repo.claim_next_job(clock.t))              # a running job cannot be claimed again
        self.assertEqual(repo.recover_stale_jobs(clock.t - 60), 0)   # heartbeat still fresh
        self.assertEqual(repo.recover_stale_jobs(clock.t + 60), 1)   # stale -> retrying
        self.assertEqual(repo.get_job(jid)["status"], "retrying")
        self.assertEqual(js.run_job(jid).status, jobs.COMPLETED)
        st = {s["stage"]: s for s in repo.list_stages(jid)}
        self.assertEqual(st["VALIDATE"]["attempt"], 1)               # completed work was kept, not redone
        self.assertEqual(repo.get_job(jid)["attempt"], 2)
        self.assertEqual(repo.get_lecture(lid)["status"], "done")

    def test_claim_is_exclusive(self):
        lid, js, jid = self._setup(segments(*SEGS))
        a = self.repo.claim_next_job(time.time())
        b = self.repo.claim_next_job(time.time())
        self.assertEqual((a["id"], b), (jid, None))

    def test_stage_outputs_persisted_for_inspection(self):
        lid, js, jid = self._setup(segments(*SEGS))
        js.run_next()
        st = self.stages(jid)
        self.assertEqual(st["VALIDATE"]["output"]["mode"], "video_slides")
        self.assertEqual(len(st["TRANSCRIBE"]["output"]["segments"]), 2)
        self.assertEqual(len(st["DETECT_SLIDE_CHANGES"]["output"]["timeline"]), 4)
        self.assertEqual(st["BUILD_EVIDENCE"]["output"]["evidence"], 2)    # PDF slides have no text layer


def jobs_validate_output(repo, lid):
    a = {x["type"]: x for x in repo.list_assets(lid)}
    return {"media_asset_id": a["video"]["id"], "media_key": a["video"]["storage_key"], "media_type": "video",
            "deck_asset_id": a["pdf"]["id"], "deck_key": a["pdf"]["storage_key"], "deck_type": "pdf",
            "mode": "video_slides"}


if __name__ == "__main__":
    unittest.main()
