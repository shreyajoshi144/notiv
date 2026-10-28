"""Format robustness evaluation: generates real-world variants of every supported input (PDF, PPTX, MP3, MP4) and of the
broken/invalid files users actually upload, runs each through the REAL validation + processing pipeline (ffmpeg,
poppler, LibreOffice, faster-whisper), and scores the share that end in the CORRECT outcome:

  valid input   -> processed to "done" with usable evidence (text / transcript)
  invalid input -> rejected at upload or failed with a clear, user-readable reason (never a crash or a silent "done")

    PYTHONPATH=.:core:backend .venv/bin/python tools/eval_format_matrix.py [--quick] [--json out.json]

No AI provider is needed (study generation is not part of this matrix; see tools/e2e_*.py). Speech is generated with macOS
`say` so transcription can be checked against the script. Synthetic files are a proxy for real uploads: the score is a
measurement on THIS matrix, not a guarantee for every file in the world.
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [ROOT, os.path.join(ROOT, "core"), os.path.join(ROOT, "backend")]
os.environ.setdefault("NOTIV_AUTH_SECRET", "matrix-" + "x" * 40)
os.environ.setdefault("NOTIV_AI_STUDY", "off")
os.environ.setdefault("NOTIV_VISION", "off")

SPEECH = ("Today we study caching. A cache keeps frequently used data in memory so the database answers fewer requests. "
          "A write through cache updates the cache and the database together.")
KEYWORDS = ("cache", "database", "memory")


def sh(*args):
    subprocess.run(args, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)


def ff(*args):
    sh("ffmpeg", "-y", "-loglevel", "error", *args)


# ── generators ───────────────────────────────────────────────────────────
def speech_wav(d):
    p = os.path.join(d, "speech.aiff")
    if not os.path.exists(p):
        sh("say", "-r", "175", "-o", p, SPEECH)
    return p


def pdf_reportlab(path, pages, rotate=None, sizes=None, image_only=False):
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import Image, ImageDraw, ImageFont
    c = canvas.Canvas(path)
    for i, lines in enumerate(pages):
        w, h = (sizes[i] if sizes else (612, 792))
        c.setPageSize((w, h))
        if image_only:
            im = Image.new("RGB", (1200, int(1200 * h / w)), "white")
            dr = ImageDraw.Draw(im)
            f = ImageFont.load_default(size=40)
            for n, ln in enumerate(lines):
                dr.text((60, 60 + n * 70), ln, fill="black", font=f)
            c.drawImage(ImageReader(im), 0, 0, w, h)
        else:
            from reportlab.lib.utils import simpleSplit
            t = c.beginText(40, h - 60)
            t.setFont("Helvetica", 14)
            for ln in lines:                                  # wrap to the page like a real document
                for part in simpleSplit(ln, "Helvetica", 14, w - 80) or [""]:
                    t.textLine(part)
            c.drawText(t)
        c.showPage()
    c.save()
    if rotate:
        from pypdf import PdfReader, PdfWriter
        r, wtr = PdfReader(path), PdfWriter()
        for n, pg in enumerate(r.pages):
            if n in rotate:
                pg.rotate(90)
            wtr.add_page(pg)
        with open(path, "wb") as f:
            wtr.write(f)
    return path


def pptx_make(path, slides, hidden=(), table=False, image_only=False, group=False):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    for i, (title, body) in enumerate(slides):
        s = prs.slides.add_slide(prs.slide_layouts[5 if (table or image_only or group) else 1])
        if not image_only and s.shapes.title is not None:
            s.shapes.title.text = title
        if table:
            tb = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
            for r, row in enumerate((("Term", "Meaning"), ("Cache", "fast memory layer"), ("Shard", "a partition of data"))):
                for col, v in enumerate(row):
                    tb.cell(r, col).text = v
        elif image_only:
            from PIL import Image, ImageDraw, ImageFont
            img = os.path.join(os.path.dirname(path), f"slide{i}.png")
            im = Image.new("RGB", (1280, 720), "white")
            ImageDraw.Draw(im).text((60, 60), title + " — " + body, fill="black", font=ImageFont.load_default(size=44))
            im.save(img)
            s.shapes.add_picture(img, Inches(0.5), Inches(0.5), Inches(9), Inches(5))
        elif group:
            g = s.shapes.add_group_shape()
            tx = g.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
            tx.text_frame.text = body
        else:
            s.placeholders[1].text = body
        if i in hidden:
            s._element.set("show", "0")
    prs.save(path)
    return path


def make_cases(d, quick):
    sp = speech_wav(d)
    P = lambda n: os.path.join(d, n)                                                  # noqa: E731
    deck = [["Caching", "A cache stores hot data in memory."], ["Sharding", "Sharding splits a table across machines."],
            ["CAP theorem", "Pick two of consistency, availability, partition tolerance."]]
    prose = [["System design notes", "Load balancers spread requests across servers so that no single server is overloaded. "
              "Replication keeps copies of data on several machines to survive failures."]]
    cases = []

    def case(name, fmt, path_fn, expect="done", needs=None):
        cases.append({"name": name, "fmt": fmt, "build": path_fn, "expect": expect, "needs": needs or []})

    # PDF
    case("pdf: text slide deck", "pdf", lambda: pdf_reportlab(P("deck.pdf"), deck), needs=["cache"])
    case("pdf: prose document", "pdf", lambda: pdf_reportlab(P("prose.pdf"), prose), needs=["replication"])
    case("pdf: scanned (image only, OCR)", "pdf", lambda: pdf_reportlab(P("scan.pdf"), deck, image_only=True), needs=["cache"])
    case("pdf: rotated pages", "pdf", lambda: pdf_reportlab(P("rot.pdf"), deck, rotate={1}), needs=["sharding"])
    case("pdf: mixed page sizes", "pdf", lambda: pdf_reportlab(P("sizes.pdf"), deck, sizes=[(612, 792), (842, 595), (420, 595)]), needs=["cap"])
    case("pdf: unicode text", "pdf", lambda: pdf_reportlab(P("uni.pdf"), [["Café résumé naïve", "Grüße — déjà vu, coöperate"]]), needs=["café"])
    if not quick:
        case("pdf: 60 pages", "pdf", lambda: pdf_reportlab(P("long.pdf"), [[f"Topic {i}", f"Point {i} about caching."] for i in range(60)]), needs=["topic 59"])
    # invalid PDFs
    case("pdf: password-protected", "pdf", lambda: _encrypted(P("enc.pdf"), deck), expect="reject")
    case("pdf: truncated / damaged", "pdf", lambda: _truncate(pdf_reportlab(P("cut.pdf"), deck), P("cut2.pdf")), expect="reject")
    case("pdf: PNG renamed .pdf", "pdf", lambda: _png_as(P("fake.pdf")), expect="reject")
    case("pdf: empty file", "pdf", lambda: _empty(P("empty.pdf")), expect="reject")
    # PPTX
    case("pptx: title + bullet slides", "pptx", lambda: pptx_make(P("deck.pptx"), [(s[0], s[1]) for s in deck]), needs=["sharding"])
    case("pptx: table slide", "pptx", lambda: pptx_make(P("table.pptx"), [("Glossary", "")], table=True), needs=["partition"])
    case("pptx: hidden slide kept in order", "pptx", lambda: pptx_make(P("hidden.pptx"), [(s[0], s[1]) for s in deck], hidden={1}), needs=["cap"])
    case("pptx: grouped shapes", "pptx", lambda: pptx_make(P("group.pptx"), [("Group", "Replication copies data to other nodes.")], group=True), needs=["replication"])
    case("pptx: image-only slides (OCR)", "pptx", lambda: pptx_make(P("img.pptx"), [(s[0], s[1]) for s in deck], image_only=True), needs=["cache"])
    if not quick:
        case("pptx: 30 slides", "pptx", lambda: pptx_make(P("long.pptx"), [(f"Slide {i}", f"Point {i}") for i in range(30)]), needs=["point 29"])
    case("pptx: zip that is not a presentation", "pptx", lambda: _zip_as(P("notppt.pptx")), expect="reject")
    # MP3
    case("mp3: 128k CBR stereo", "mp3", lambda: _enc(sp, P("cbr.mp3"), "-ac", "2", "-b:a", "128k"), needs=list(KEYWORDS))
    case("mp3: VBR", "mp3", lambda: _enc(sp, P("vbr.mp3"), "-q:a", "4"), needs=list(KEYWORDS))
    case("mp3: 8 kHz mono (phone quality)", "mp3", lambda: _enc(sp, P("8k.mp3"), "-ac", "1", "-ar", "8000", "-b:a", "16k"), needs=["cache"])
    case("mp3: with cover art", "mp3", lambda: _cover(sp, P("cover.mp3"), d), needs=list(KEYWORDS))
    case("mp3: speech with background noise", "mp3", lambda: _noisy(sp, P("noise.mp3")), needs=["cache"])
    case("mp3: silence only", "mp3", lambda: _silence(P("silence.mp3")), expect="fail_clear")
    case("mp3: WAV renamed .mp3", "mp3", lambda: _wav_as(sp, P("wav.mp3")), expect="reject")
    # MP4
    case("mp4: H.264 + AAC", "mp4", lambda: _video(sp, P("h264.mp4")), needs=list(KEYWORDS))
    case("mp4: H.264 + MP3 audio", "mp4", lambda: _video(sp, P("mp3a.mp4"), acodec="libmp3lame"), needs=list(KEYWORDS))
    case("mp4: HEVC + AAC", "mp4", lambda: _video(sp, P("hevc.mp4"), vcodec="libx265"), needs=list(KEYWORDS))
    case("mp4: vertical 9:16 at 60 fps", "mp4", lambda: _video(sp, P("vert.mp4"), size="360x640", rate=60), needs=list(KEYWORDS))
    case("mp4: picture longer than sound", "mp4", lambda: _video(sp, P("longpic.mp4"), extra_seconds=6), needs=list(KEYWORDS))
    case("mp4: sound longer than picture", "mp4", lambda: _video(sp, P("shortpic.mp4"), pic_seconds=3), needs=list(KEYWORDS))
    case("mp4: no audio track", "mp4", lambda: _video(None, P("mute.mp4")), expect="reject")
    case("mp4: MKV renamed .mp4", "mp4", lambda: _mkv_as(sp, P("mkv.mp4")), expect="reject")
    return cases


def _encrypted(path, deck):
    from pypdf import PdfReader, PdfWriter
    src = pdf_reportlab(path + ".src.pdf", deck)
    w = PdfWriter(clone_from=PdfReader(src))
    w.encrypt("secret-pass")
    with open(path, "wb") as f:
        w.write(f)
    return path


def _truncate(src, dst):
    data = open(src, "rb").read()
    open(dst, "wb").write(data[: len(data) // 3])
    return dst


def _png_as(path):
    from PIL import Image
    Image.new("RGB", (40, 40), "red").save(path, format="PNG")
    return path


def _empty(path):
    open(path, "wb").close()
    return path


def _zip_as(path):
    import zipfile
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("readme.txt", "not a presentation")
    return path


def _enc(src, out, *args):
    ff("-i", src, *args, "-c:a", "libmp3lame", out)
    return out


def _cover(src, out, d):
    jpg = os.path.join(d, "cover.jpg")
    ff("-f", "lavfi", "-i", "color=c=red:s=300x300:d=1", "-frames:v", "1", jpg)
    ff("-i", src, "-i", jpg, "-map", "0:a", "-map", "1:v", "-c:a", "libmp3lame", "-c:v", "mjpeg", "-id3v2_version", "3",
       "-disposition:v", "attached_pic", out)
    return out


def _noisy(src, out):
    ff("-i", src, "-f", "lavfi", "-i", "anoisesrc=color=pink:amplitude=0.05:d=20", "-filter_complex",
       "[0:a][1:a]amix=inputs=2:duration=first", "-c:a", "libmp3lame", out)
    return out


def _silence(out):
    ff("-f", "lavfi", "-i", "anullsrc=r=44100:cl=mono", "-t", "8", "-c:a", "libmp3lame", out)
    return out


def _wav_as(src, out):
    ff("-i", src, "-f", "wav", out)
    return out


def _video(src, out, vcodec="libx264", acodec="aac", size="640x360", rate=10, extra_seconds=0, pic_seconds=None):
    dur = _dur(src) if src else 8
    pic = pic_seconds if pic_seconds is not None else dur + extra_seconds
    args = ["-f", "lavfi", "-i", f"testsrc=size={size}:rate={rate}:duration={pic}"]
    if src:
        args += ["-i", src, "-c:a", acodec]
    args += ["-c:v", vcodec, "-pix_fmt", "yuv420p"]
    if vcodec == "libx265":
        args += ["-tag:v", "hvc1"]
    ff(*args, out)
    return out


def _mkv_as(src, out):
    ff("-f", "lavfi", "-i", "testsrc=size=320x240:rate=10:duration=3", "-i", src, "-c:v", "libx264", "-c:a", "aac", "-f", "matroska", out)
    return out


def _dur(p):
    r = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", p], capture_output=True, text=True)
    return float(r.stdout.strip() or 0)


# ── runner ───────────────────────────────────────────────────────────────
CLEAR = re.compile(r"(not a|unsupported|password|damaged|unreadable|no audio|no speech|empty|not valid|bad signature|"
                   r"supported formats|not a valid|contains no|no pages)", re.I)


def run_case(c, work):
    from core.lecture import validation
    from core.lecture.errors import ValidationError
    from core.lecture.testkit import Env
    t0 = time.monotonic()
    path = c["build"]()
    env = Env(tempfile.mkdtemp(dir=work))
    is_media = c["fmt"] in ("mp3", "mp4")
    try:
        lid = env.lecture(media=path if is_media else None, deck=None if is_media else path)
    except ValidationError as e:
        ok = c["expect"] == "reject" and bool(CLEAR.search(str(e)))
        return {"outcome": "rejected", "ok": ok, "detail": str(e)[:160], "s": round(time.monotonic() - t0, 1)}
    except Exception as e:                                               # noqa: BLE001 — a crash is a failure of the matrix
        return {"outcome": "crash_at_upload", "ok": False, "detail": f"{type(e).__name__}: {e}"[:200], "s": round(time.monotonic() - t0, 1)}
    try:
        env.process(lid)
    except Exception as e:                                               # noqa: BLE001
        lec = env.repo.get_lecture(lid)
        clear = bool(CLEAR.search(lec.get("error") or str(e)))
        ok = c["expect"] == "fail_clear" and lec["status"] == "failed" and clear
        return {"outcome": "failed", "ok": ok, "detail": (lec.get("error") or str(e))[:200], "s": round(time.monotonic() - t0, 1)}
    ev = env.repo.list_evidence(lid)
    text = " ".join(e["text"] for e in ev).lower()
    missing = [k for k in c["needs"] if k.lower() not in text]
    ok = c["expect"] == "done" and bool(ev) and not missing
    return {"outcome": "done", "ok": ok, "evidence": len(ev), "missing": missing, "s": round(time.monotonic() - t0, 1),
            "detail": "" if ok else (f"missing words: {missing}" if missing else f"expected {c['expect']}")}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--quick", action="store_true")
    ap.add_argument("--json")
    a = ap.parse_args()
    work = tempfile.mkdtemp(prefix="notiv-matrix-")
    os.environ["NOTIV_LECTURE_DB"] = os.path.join(work, "unused.db")
    results = []
    try:
        cases = make_cases(work, a.quick)
        print(f"{len(cases)} inputs\n")
        for c in cases:
            try:
                r = run_case(c, work)
            except Exception as e:                                       # noqa: BLE001 — generator problems are reported, not scored
                r = {"outcome": "generator_error", "ok": None, "detail": f"{type(e).__name__}: {e}"[:200]}
                traceback.print_exc()
            results.append({"name": c["name"], "fmt": c["fmt"], "expect": c["expect"], **r})
            mark = "PASS" if r["ok"] else ("SKIP" if r["ok"] is None else "FAIL")
            print(f"[{mark}] {c['name']:<40} -> {r['outcome']:<10} {r.get('s', '')}s  {r.get('detail', '')}")
    finally:
        shutil.rmtree(work, ignore_errors=True)
    scored = [r for r in results if r["ok"] is not None]
    passed = sum(1 for r in scored if r["ok"])
    by = {}
    for r in scored:
        by.setdefault(r["fmt"], [0, 0])
        by[r["fmt"]][0] += int(r["ok"])
        by[r["fmt"]][1] += 1
    print("\nby format: " + ", ".join(f"{k.upper()} {v[0]}/{v[1]}" for k, v in by.items()))
    print(f"correct outcome: {passed}/{len(scored)} = {100 * passed / max(1, len(scored)):.1f}%")
    if a.json:
        json.dump({"results": results, "passed": passed, "total": len(scored)}, open(a.json, "w"), indent=1)
    return 0 if passed == len(scored) else 1


if __name__ == "__main__":
    sys.exit(main())
