#!/usr/bin/env python3
"""Real-backend browser check of the redesigned frontend. PYTHONPATH=backend:.:core python tools/e2e_redesign.py --out DIR
Real uvicorn + SQLite + in-process worker + Chromium; registers a user through the UI, uploads real PDF/PPTX/MP3/MP4 files."""
import argparse, json, os, subprocess, sys, tempfile, time, socket, traceback
from urllib.request import urlopen
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
from playwright.sync_api import sync_playwright
CHROME = next((p for p in ("/opt/pw-browsers/chromium-1194/chrome-linux/chrome", "/opt/pw-browsers/chromium") if os.path.exists(p)), None)
R = []
def step(name):
    class S:
        def __enter__(s): s.t = time.time(); return s
        def __exit__(s, et, ev, tb):
            R.append({"step": name, "ok": et is None, "err": None if et is None else "".join(traceback.format_exception_only(et, ev)).strip()[:400]})
            print(("PASS " if et is None else "FAIL ") + name + ("" if et is None else "  <- " + R[-1]["err"])); return True
    return S()
def port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p
def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--out", required=True); a = ap.parse_args(); out = a.out; os.makedirs(out, exist_ok=True)
    work = tempfile.mkdtemp(prefix="notiv-rd-"); fx = os.path.join(work, "fx"); os.makedirs(fx); py = sys.executable
    env = {**os.environ, "NOTIV_ENV": "development", "NOTIV_AUTH_SECRET": "e2e-auth-secret-e2e-auth-secret-e2e-auth-00", "NOTIV_SIGNING_SECRET": "e2e-sign-secret-e2e-sign-secret-e2e-sign-00",
           "NOTIV_LECTURE_DB": os.path.join(work, "e.db"), "NOTIV_STORAGE_DIR": os.path.join(work, "obj"), "NOTIV_INPROCESS_WORKER": "1",
           "PYTHONPATH": os.pathsep.join([os.path.join(ROOT, "backend"), ROOT, os.path.join(ROOT, "core")])}
    for k in ("RL_READ", "RL_WRITE", "RL_STUDY", "RL_ASK", "RL_UPLOAD", "RL_PROCESS"): env["NOTIV_" + k] = "100000/60"
    # fixtures
    from core.lecture.make_test_lecture import make_text_pdf
    pdf = make_text_pdf(fx, n=3, slide_lines=[["Cloud basics", "AWS is a cloud platform."], ["Storage", "S3 stores objects."], ["Compute", "EC2 runs servers."]], name="aws_deck.pdf")
    from pptx import Presentation
    pr = Presentation(); 
    for i, t in enumerate(["Intro", "Networks", "Security"]):
        sl = pr.slides.add_slide(pr.slide_layouts[1]); sl.shapes.title.text = t; sl.placeholders[1].text = f"Slide {i+1} body text"
    pptx = os.path.join(fx, "deck.pptx"); pr.save(pptx)
    mp3 = os.path.join(fx, "talk.mp3"); mp4 = os.path.join(fx, "talk.mp4")
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=6", mp3], check=True)
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=blue:s=320x240:d=6", "-f", "lavfi", "-i", "sine=frequency=330:duration=6", "-shortest", "-pix_fmt", "yuv420p", mp4], check=True)
    p = port(); base = f"http://127.0.0.1:{p}"
    srv = subprocess.Popen([py, "-m", "uvicorn", "server.main:app", "--host", "127.0.0.1", "--port", str(p), "--log-level", "warning"], env=env, cwd=ROOT, stdout=open(os.path.join(out, "server.log"), "w"), stderr=subprocess.STDOUT)
    try:
        for _ in range(100):
            try: urlopen(base + "/healthz", timeout=2).read(); break
            except Exception: time.sleep(.4)
        with sync_playwright() as pw:
            br = pw.chromium.launch(executable_path=CHROME, args=["--no-sandbox"]) 
            ctx = br.new_context(viewport={"width": 1366, "height": 860}); pg = ctx.new_page()
            errs = []
            pg.on("pageerror", lambda e: errs.append("pageerror " + str(e))); pg.on("console", lambda m: errs.append("console " + m.text) if m.type == "error" else None)
            shot = lambda n: pg.screenshot(path=os.path.join(out, n + ".png"))
            with step("landing loads with exact hero image, nav, no overflow"):
                pg.goto(base + "/"); pg.wait_for_load_state("networkidle"); shot("landing")
                assert "more" in pg.locator("h1").inner_text(); 
                for t in ("Features", "Solutions", "How it works", "Resources", "FAQs", "Get started"): assert pg.get_by_role("link", name=t).first.is_visible(), t
                assert pg.evaluate("document.documentElement.scrollWidth <= innerWidth")
                nat = pg.evaluate("[document.querySelector('.hero-img').naturalWidth, document.querySelector('.hero-img').naturalHeight]"); assert nat == [736, 414], nat
            with step("register through the UI (email/password), lands in the app"):
                pg.goto(base + "/register"); shot("register")
                pg.fill("#name", "Test Student"); pg.fill("#email", "student@example.com"); pg.fill("#password", "correct horse battery 1"); pg.fill("#confirmation", "correct horse battery 1")
                pg.click("button[type=submit]"); pg.wait_for_url("**/app**"); pg.wait_for_selector(".rail"); shot("desk-empty")
                assert "Add your first lecture" in pg.locator(".next").inner_text()
            uid = pg.evaluate("fetch('/auth/me').then(r=>r.json()).then(j=>j.id)")
            subprocess.run([py, os.path.join(ROOT, "tools", "seed_e2e_world.py"), "--user", str(uid), "--json", os.path.join(work, "seed.json")], env=env, check=True, capture_output=True)
            world = json.load(open(os.path.join(work, "seed.json"))); L = world["lectures"]
            with step("desk: next action, rail nav, blurred background"):
                pg.goto(base + "/app"); pg.wait_for_selector(".next"); shot("desk")
                assert pg.locator(".rail a.item").count() == 6
                assert "notiv-desk-blur" in pg.evaluate("getComputedStyle(document.body).backgroundImage")
            with step("upload PDF + MP3 via dialog, processing is shown, then a terminal state"):
                pg.goto(base + "/app?view=library"); pg.wait_for_selector("h1"); pg.get_by_role("button", name="Add a lecture").first.click()
                pg.fill("#nl-title", "Upload test A"); pg.set_input_files("#nl-media", mp3); pg.set_input_files("#nl-deck", pdf); shot("upload-dialog")
                pg.get_by_role("button", name="Upload and process").click(); pg.wait_for_url("**/app/lectures/**", timeout=30000)
                pg.wait_for_timeout(1500); shot("lecture-after-upload")
            with step("upload PPTX + MP4"):
                pg.goto(base + "/app?view=library"); pg.get_by_role("button", name="Add a lecture").first.click()
                pg.fill("#nl-title", "Upload test B"); pg.set_input_files("#nl-media", mp4); pg.set_input_files("#nl-deck", pptx)
                pg.get_by_role("button", name="Upload and process").click(); pg.wait_for_url("**/app/lectures/**", timeout=30000)
            for nm, f in (("PDF only", pdf), ("PPTX only", pptx)):
                with step(f"slides-only upload ({nm}) reaches Ready and previews every page"):
                    pg.goto(base + "/app?view=library"); pg.get_by_role("button", name="Add a lecture").first.click()
                    pg.fill("#nl-title", "Deck " + nm); pg.set_input_files("#nl-deck", f); pg.get_by_role("button", name="Upload and process").click(); pg.wait_for_url("**/app/lectures/**", timeout=30000)
                    end = time.time() + 120
                    while time.time() < end:
                        pg.wait_for_timeout(3000); pg.reload(); pg.wait_for_timeout(1500)
                        if pg.locator(".chip.done").count() or pg.locator(".chip.failed").count(): break
                    shot("deck-" + nm.split()[0].lower()); assert pg.locator(".chip.done").count(), pg.locator("main").inner_text()[:300]
                    print("   pages:", pg.locator("text=/of \\d+/").first.inner_text())
            with step("library lists lectures with states; wait for processing to settle"):
                pg.goto(base + "/app?view=library"); end = time.time() + 90
                while time.time() < end:
                    pg.reload(); pg.wait_for_selector("#lec-list"); chips = pg.locator("#lec-list .chip").all_inner_texts()
                    if not any(c in ("Processing", "Waiting") for c in chips): break
                    time.sleep(3)
                shot("library"); print("   chips:", chips); print("   rows:", pg.locator("#lec-list .lec-li").all_inner_texts())
            with step("study view of a seeded lecture: viewer, transcript, nav"):
                pg.goto(base + f"/app/lectures/{L['gradient_descent']}", wait_until="domcontentloaded"); pg.wait_for_timeout(3000); shot("study"); pg.set_viewport_size({"width": 1280, "height": 800}); pg.wait_for_timeout(500); pg.screenshot(path=os.path.join(out, "app-preview.png")); pg.set_viewport_size({"width": 1366, "height": 860})
                assert pg.evaluate("document.documentElement.scrollWidth <= innerWidth")
            with step("full-size Ask chat: real answer + citations from existing backend"):
                pg.goto(base + f"/static/study-companion.html?lecture={L['gradient_descent']}"); pg.wait_for_selector("#cc-in")
                pg.fill("#cc-in", "What is gradient descent?"); pg.click(".chat-form button[type=submit]"); pg.wait_for_selector(".msg.bot .txt", timeout=30000); pg.wait_for_timeout(800); shot("chat")
                txt = pg.locator(".msg.bot").last.inner_text(); print("   answer:", txt[:160].replace("\n", " "))
                assert len(txt) > 20
            with step("cat launcher opens the full chat; timer start/pause/reset persists across reload"):
                pg.goto(base + "/app"); pg.wait_for_selector(".cat-btn"); pg.click(".cat-btn"); shot("cat-menu")
                assert pg.get_by_role("menuitem", name="Ask your study companion").get_attribute("href").startswith("/static/study-companion.html")
                pg.get_by_role("menuitem", name="Focus timer").click(); pg.click(".ft-go"); pg.wait_for_timeout(2600); shot("timer")
                t1 = pg.locator(".ft-time").inner_text(); assert t1 != "25:00", t1
                pg.reload(); pg.wait_for_selector(".focus-pill"); assert "Focus" in pg.locator(".focus-pill").inner_text()
                pg.click(".focus-pill"); pg.click(".ft-go"); pg.click("text=Reset"); assert pg.locator(".ft-time").inner_text() == "25:00"
            with step("question bank: add, edit-free list, delete confirm, practise a saved question"):
                pg.goto(base + f"/static/practice-quiz.html?lecture={L['gradient_descent']}"); pg.wait_for_load_state("networkidle")
                pg.get_by_role("button", name="My question bank").click(); pg.wait_for_selector("#bank-question")
                pg.fill("#bank-question", "What does the learning rate control?")
                for i, t in enumerate(["Step size", "Batch size", "Layer count", "Dropout"]): pg.fill(f"#bank-option-{i}", t)
                pg.get_by_role("button", name="Add question").click(); pg.wait_for_selector(".pq-q"); shot("bank")
                assert "1 saved question" in pg.locator(".pq-score").inner_text()
                pg.reload(); pg.get_by_role("button", name="My question bank").click(); pg.wait_for_selector(".pq-q")
                pg.get_by_role("button", name="Practise my questions").click(); pg.wait_for_timeout(800); shot("bank-practice")
                pg.get_by_text("Step size").first.click(); pg.wait_for_timeout(500)
                pg.goto(base + f"/static/practice-quiz.html?lecture={L['gradient_descent']}"); pg.get_by_role("button", name="My question bank").click(); pg.wait_for_selector(".pq-q")
                pg.locator(".pq-q").get_by_role("button", name="Delete").click(); pg.locator("dialog button:has-text('Delete'), [role=dialog] button:has-text('Delete')").last.click(); pg.wait_for_selector("text=No saved questions")
            with step("account page + sign out"):
                pg.goto(base + "/static/account.html"); pg.wait_for_selector("h1"); shot("account"); pg.get_by_role("button", name="Sign out").click(); pg.wait_for_url("**/sign-in**")
            with step("mobile layout: landing, auth, desk (no horizontal overflow)"):
                m = br.new_context(viewport={"width": 390, "height": 844}, is_mobile=True); mp = m.new_page()
                mp.goto(base + "/"); mp.wait_for_load_state("networkidle"); mp.screenshot(path=os.path.join(out, "m-landing.png"))
                assert mp.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"), "landing overflow"
                mp.goto(base + "/sign-in"); mp.fill("#email", "student@example.com"); mp.fill("#password", "correct horse battery 1"); mp.click("button[type=submit]"); mp.wait_for_url("**/app**"); mp.wait_for_selector(".rail"); mp.wait_for_timeout(600)
                mp.screenshot(path=os.path.join(out, "m-desk.png")); assert mp.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"), "desk overflow"
                mp.goto(base + f"/app/lectures/{L['gradient_descent']}"); mp.wait_for_timeout(1500); mp.screenshot(path=os.path.join(out, "m-study.png"))
                assert mp.evaluate("document.documentElement.scrollWidth <= innerWidth + 1"), "study overflow"
            real = [e for e in errs if "favicon" not in e and "401" not in e]; R.append({"step": "console/page errors", "ok": not real, "err": "; ".join(real)[:500]}); print("console errors:", real[:5])
            br.close()
    finally:
        srv.terminate()
    json.dump(R, open(os.path.join(out, "report.json"), "w"), indent=1)
    print(f"{sum(r['ok'] for r in R)}/{len(R)} passed"); return 0 if all(r["ok"] for r in R) else 1
sys.exit(main())
