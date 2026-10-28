"""Notebook API: persistent, owner-scoped content, canvas, exports and revocable read-only shares."""
import csv, io, json, re, zipfile, html
from typing import Any, Dict, Optional
from fastapi import APIRouter, Depends, Body, Query, HTTPException
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field
from server.api.deps import get_current_user
from server.api.workspace_handlers import ApiError
from core.lecture.errors import NotFoundError, ValidationError
from server.api import notebooks_service as svc

router=APIRouter(tags=["notebooks"])
def run(fn,*a,**kw):
    try:return fn(*a,**kw)
    except NotFoundError as e: raise HTTPException(404,"not found")
    except ValidationError as e: raise HTTPException(409 if "Version conflict" in str(e) or "changed elsewhere" in str(e) else 400,str(e))

class NotebookIn(BaseModel):
    title: str = Field(..., min_length=1, max_length=200)
    description: str = Field("", max_length=2000)
    subject: str = Field("", max_length=120)
    workspace_id: Optional[int] = Field(None, gt=0)
    model_config = {"extra": "forbid"}

class SourceIn(BaseModel):
    source_type: str = "lecture"
    source_id: int = Field(..., gt=0)
    label: str = Field("", max_length=300)
    metadata: Dict[str, Any] = Field(default_factory=dict)
    model_config = {"extra": "forbid"}

class ItemIn(BaseModel):
    item_type: str = Field(..., max_length=40)
    title: str = Field(..., min_length=1, max_length=240)
    content: Any
    provenance: list = Field(default_factory=list)
    expected_version: Optional[int] = Field(None, ge=1)
    model_config = {"extra": "forbid"}

class CanvasIn(BaseModel):
    state: Dict[str, Any]
    expected_version: Optional[int] = Field(None, ge=1)
    model_config = {"extra": "forbid"}

class ShareIn(BaseModel):
    permission: str = "read"
    expires_in_seconds: Optional[int] = None
    model_config = {"extra": "forbid"}

@router.get("/notebooks")
def notebooks(q:str=Query("",max_length=120),include_archived:bool=False,user:dict=Depends(get_current_user)):return run(svc.list_notebooks,user["id"],q,include_archived)
@router.post("/notebooks",status_code=201)
def notebook_create(body:NotebookIn,user:dict=Depends(get_current_user)):return run(svc.create_notebook,user["id"],body.title,body.description,body.subject,body.workspace_id)
@router.get("/notebooks/{notebook_id}")
def notebook_get(notebook_id:int,user:dict=Depends(get_current_user)):return run(svc.get_notebook,user["id"],notebook_id)
@router.patch("/notebooks/{notebook_id}")
def notebook_patch(notebook_id:int,body:Dict[str,Any]=Body(...),user:dict=Depends(get_current_user)):
 if set(body)-{"title","description","subject","status","workspace_id"}:raise HTTPException(422,"Unsupported fields")
 return run(svc.update_notebook,user["id"],notebook_id,body)
@router.delete("/notebooks/{notebook_id}",status_code=204)
def notebook_delete(notebook_id:int,user:dict=Depends(get_current_user)):run(svc.delete_notebook,user["id"],notebook_id);return Response(status_code=204)
@router.post("/notebooks/{notebook_id}/sources",status_code=201)
def source_add(notebook_id:int,body:SourceIn,user:dict=Depends(get_current_user)):return run(svc.add_source,user["id"],notebook_id,body.source_type,body.source_id,body.label,body.metadata)
@router.delete("/notebooks/{notebook_id}/sources/{source_id}",status_code=204)
def source_delete(notebook_id:int,source_id:int,user:dict=Depends(get_current_user)):run(svc.remove_source,user["id"],notebook_id,source_id);return Response(status_code=204)
@router.post("/notebooks/{notebook_id}/items",status_code=201)
def item_create(notebook_id:int,body:ItemIn,user:dict=Depends(get_current_user)):return run(svc.save_item,user["id"],notebook_id,None,body.item_type,body.title,body.content,body.provenance)
@router.put("/notebooks/{notebook_id}/items/{item_id}")
def item_update(notebook_id:int,item_id:int,body:ItemIn,user:dict=Depends(get_current_user)):return run(svc.save_item,user["id"],notebook_id,item_id,body.item_type,body.title,body.content,body.provenance,body.expected_version)
@router.delete("/notebooks/{notebook_id}/items/{item_id}",status_code=204)
def item_delete(notebook_id:int,item_id:int,user:dict=Depends(get_current_user)):run(svc.delete_item,user["id"],notebook_id,item_id);return Response(status_code=204)
@router.put("/notebooks/{notebook_id}/canvas")
def canvas_save(notebook_id:int,body:CanvasIn,user:dict=Depends(get_current_user)):return run(svc.save_canvas,user["id"],notebook_id,body.state,body.expected_version)
@router.get("/notebooks/{notebook_id}/shares")
def shares_list(notebook_id:int,user:dict=Depends(get_current_user)):return run(svc.list_shares,user["id"],notebook_id)
@router.post("/notebooks/{notebook_id}/shares",status_code=201)
def share_create(notebook_id:int,body:ShareIn,user:dict=Depends(get_current_user)):return run(svc.create_share,user["id"],notebook_id,body.permission,body.expires_in_seconds)
@router.delete("/notebooks/{notebook_id}/shares/{share_id}",status_code=204)
def share_revoke(notebook_id:int,share_id:int,user:dict=Depends(get_current_user)):run(svc.revoke_share,user["id"],notebook_id,share_id);return Response(status_code=204)
@router.get("/shared/notebooks/{token}")
def share_view(token:str):
 if len(token)>200:raise HTTPException(404,"share not found or expired")
 return run(svc.public_share,token)


class GenerateIn(BaseModel):
    source_id: int = Field(..., gt=0)
    output_type: str = Field(..., pattern="^(summary|concepts|definitions|flashcards|quiz|mind_map)$")
    title: Optional[str] = Field(None, max_length=240)
    model_config = {"extra": "forbid"}

@router.post("/notebooks/{notebook_id}/generate", status_code=201)
def generate_notebook_item(notebook_id: int, body: GenerateIn, user: dict = Depends(get_current_user)):
    # Generation is delegated to the existing evidence-grounded lecture/study services. Only linked sources are eligible.
    nb = run(svc.get_notebook, user["id"], notebook_id)
    linked = any(s.get("source_type") == "lecture" and int(s.get("source_id")) == body.source_id for s in nb.get("sources", []))
    if not linked: raise HTTPException(400, "Link this lecture to the notebook before generating from it")
    ai_status = None
    try:
        from server.api.workspace_handlers import ensure_ai_study
        ai_status = ensure_ai_study(user["id"], body.source_id)       # AI study material first (no-op when current)
        if body.output_type == "summary":
            from server.api.workspace_handlers import summary
            content = summary(user["id"], body.source_id)
        elif body.output_type in ("concepts", "definitions"):
            from server.api.study_handlers import knowledge
            content = knowledge(user["id"], body.source_id, "concept" if body.output_type == "concepts" else "definition")
        elif body.output_type == "mind_map":
            # topic and concept nodes with labelled edges from the AI study material; never a map without connections
            from server.api.workspace_handlers import mind_map
            content = mind_map(user["id"], body.source_id)
            if not content.get("nodes") and ai_status and ai_status.get("message"):
                content["message"] = ai_status["message"]
        else:
            from server.api.study_handlers import study_generate, study_items
            kinds = ["flashcard"] if body.output_type == "flashcards" else ["mcq"]
            study_generate(user["id"], body.source_id, kinds)
            content = study_items(user["id"], body.source_id, kind=kinds[0], reveal=True)
    except Exception as exc:
        if isinstance(exc, HTTPException): raise
        if hasattr(exc, "status"): raise HTTPException(exc.status, getattr(exc, "detail", "Generation failed"))
        raise
    if isinstance(content, dict) and ai_status and ai_status.get("status") in ("failed", "no_provider") and content.get("mode") != "ai":
        # say WHY this is verbatim / empty (no credits, rate limit, bad key, ...) instead of implying the lecture had nothing
        content = {**content, "ai_unavailable": ai_status.get("message")}
    title = body.title or (body.output_type.replace("_", " ").title() + " · Lecture " + str(body.source_id))
    ai_used = isinstance(content, dict) and (content.get("mode") == "ai" or (ai_status or {}).get("status") in ("ok", "cached"))
    provenance = [{"source_type":"lecture", "source_id":body.source_id, "output_type":body.output_type,
                   "generated_by": "ai_study" if ai_used else "existing_notiv_grounded_service",
                   **({"ai_status": {k: ai_status[k] for k in ("status", "message", "provider", "model") if k in ai_status}} if ai_status else {})}]
    return run(svc.save_item, user["id"], notebook_id, None, body.output_type, title, content, provenance)

@router.get("/notebooks/{notebook_id}/items/{item_id}/export")
def export_item(notebook_id: int, item_id: int, format: str = Query("md", pattern="^(json|md|txt|csv)$"), user: dict = Depends(get_current_user)):
    nb = run(svc.get_notebook, user["id"], notebook_id)
    item = next((x for x in nb.get("items", []) if int(x.get("id", 0)) == item_id), None)
    if not item: raise HTTPException(404, "not found")
    base = re.sub(r"[^A-Za-z0-9._-]+", "-", item["title"]).strip("-._")[:70] or "notiv-item"
    content = item.get("content")
    if format == "json": data=json.dumps(item,ensure_ascii=False,indent=2).encode();media="application/json; charset=utf-8"
    elif format in ("md","txt"):
        text = "# "+item["title"]+"\n\n"+(content if isinstance(content,str) else json.dumps(content,ensure_ascii=False,indent=2))+"\n"
        if item.get("provenance"): text += "\nSource references: "+json.dumps(item["provenance"],ensure_ascii=False)+"\n"
        data=text.encode();media="text/markdown; charset=utf-8" if format=="md" else "text/plain; charset=utf-8"
    else:
        buf=io.StringIO(newline="");w=csv.writer(buf);w.writerow(["title","type","content","provenance"]);prov=json.dumps(item.get("provenance"),ensure_ascii=False)
        for entry in (_csv_entries(item["item_type"],content) or [content]):w.writerow([item["title"],item["item_type"],json.dumps(entry,ensure_ascii=False),prov])
        data=buf.getvalue().encode("utf-8-sig");media="text/csv; charset=utf-8"
    return Response(data,media_type=media,headers={"Content-Disposition":f'attachment; filename="{base}.{format}"',"X-Content-Type-Options":"nosniff","Cache-Control":"no-store"})

def _csv_entries(item_type,content):
 """One CSV row per card/question. Real content is {"items":[...]}; a bare list (older saves) is accepted too."""
 if item_type not in ("flashcards","quiz"):return None
 if isinstance(content,dict) and isinstance(content.get("items"),list):return content["items"]
 if isinstance(content,list):return content
 return None
def _filename(title,ext):
 base=re.sub(r"[^A-Za-z0-9._-]+","-",title).strip("-._")[:70] or "notiv-notebook"
 return f"{base}.{ext}"
def _markdown(nb):
 out=[f"# {nb['title']}",""]
 if nb.get("subject"):out += [f"**Subject:** {nb['subject']}",""]
 if nb.get("description"):out += [nb["description"],""]
 if nb.get("sources"):
  out += ["## Sources",""]+[f"- {s.get('label') or s['source_type']+' '+str(s['source_id'])}" for s in nb["sources"]]+[""]
 for item in nb.get("items",[]):
  out += ["## "+item["title"],""]
  content=item.get("content")
  out += [content if isinstance(content,str) else "```json\n"+json.dumps(content,ensure_ascii=False,indent=2)+"\n```",""]
  prov=item.get("provenance") or []
  if prov:out += ["**Source references**",""]+[f"- {json.dumps(x,ensure_ascii=False)}" for x in prov]+[""]
 return "\n".join(out).strip()+"\n"
@router.get("/notebooks/{notebook_id}/export")
def export_notebook(notebook_id:int,format:str=Query("json",pattern="^(json|md|txt|csv|docx|pdf|svg|pptx)$"),user:dict=Depends(get_current_user)):
 nb=run(svc.get_notebook,user["id"],notebook_id);name=_filename(nb["title"],format);media="application/octet-stream";data=b""
 if format=="json":media="application/json; charset=utf-8";data=json.dumps(nb,ensure_ascii=False,indent=2).encode()
 elif format in ("md","txt"):media="text/markdown; charset=utf-8" if format=="md" else "text/plain; charset=utf-8";data=_markdown(nb).encode()
 elif format=="csv":
  buf=io.StringIO(newline="");w=csv.writer(buf);w.writerow(["title","type","content","provenance"])
  for i in nb.get("items",[]):
   content=i.get("content")
   entries=_csv_entries(i.get("item_type"),content)
   if entries is not None:
    for entry in entries:w.writerow([i["title"],i["item_type"],json.dumps(entry,ensure_ascii=False),json.dumps(i.get("provenance"),ensure_ascii=False)])
   else:w.writerow([i["title"],i["item_type"],json.dumps(content,ensure_ascii=False),json.dumps(i.get("provenance"),ensure_ascii=False)])
  media="text/csv; charset=utf-8";data=buf.getvalue().encode("utf-8-sig")
 elif format=="svg":
  state=nb.get("canvas") or {};nodes=state.get("nodes",[]);width=1200;height=max(400,100+len(nodes)*120);parts=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#f5efe6"/><text x="40" y="45" font-size="26" fill="#292622">{html.escape(nb["title"])}</text>']
  positions={}
  for ix,n in enumerate(nodes):
   x=max(20,min(900,float(n.get("x",40+(ix%3)*360))));y=max(70,min(height-90,float(n.get("y",80+(ix//3)*120))));positions[str(n.get("id"))]=(x,y)
  for edge in state.get("edges",[]):
   a=positions.get(str(edge.get("from")));b=positions.get(str(edge.get("to")))
   if a and b:parts.append(f'<line x1="{a[0]+130}" y1="{a[1]+36}" x2="{b[0]+130}" y2="{b[1]+36}" stroke="#a95e37" stroke-width="2"/>')
  for ix,n in enumerate(nodes):
   x,y=positions[str(n.get("id"))];parts.append(f'<rect x="{x}" y="{y}" width="260" height="72" rx="8" fill="#fffaf3" stroke="#a95e37"/><text x="{x+12}" y="{y+28}" font-size="15" fill="#292622">{html.escape(str(n.get("text",""))[:80])}</text>')
  parts.append('</svg>');media="image/svg+xml";data="".join(parts).encode()
 elif format in ("docx","pdf","pptx"):
  try:
   if format=="docx":
    from docx import Document
    d=Document();d.add_heading(nb["title"],0)
    for item in nb.get("items",[]):
     d.add_heading(item["title"],1);d.add_paragraph(item.get("content") if isinstance(item.get("content"),str) else json.dumps(item.get("content"),ensure_ascii=False,indent=2))
     if item.get("provenance"):d.add_paragraph("Source references: "+json.dumps(item["provenance"],ensure_ascii=False))
    buf=io.BytesIO();d.save(buf);data=buf.getvalue();media="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
   elif format=="pdf":
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    buf=io.BytesIO();doc=SimpleDocTemplate(buf);styles=getSampleStyleSheet();story=[Paragraph(html.escape(nb["title"]),styles["Title"]),Spacer(1,14)]
    for item in nb.get("items",[]):
     story += [Paragraph(html.escape(item["title"]),styles["Heading2"]),Paragraph(html.escape(item.get("content") if isinstance(item.get("content"),str) else json.dumps(item.get("content"),ensure_ascii=False)),styles["BodyText"])]
     if item.get("provenance"):story.append(Paragraph("Source references: "+html.escape(json.dumps(item["provenance"],ensure_ascii=False)),styles["BodyText"]))
     story.append(Spacer(1,8))
    doc.build(story);data=buf.getvalue();media="application/pdf"
   else:
    from pptx import Presentation
    from pptx.util import Inches,Pt
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[0]);slide.shapes.title.text=nb["title"];slide.placeholders[1].text=nb.get("description","")
    for item in nb.get("items",[]):
     slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=item["title"];tf=slide.placeholders[1].text_frame;tf.text=item.get("content") if isinstance(item.get("content"),str) else json.dumps(item.get("content"),ensure_ascii=False,indent=2)
     if item.get("provenance"):tf.text += "\n\nSource references: "+json.dumps(item["provenance"],ensure_ascii=False)
    buf=io.BytesIO();prs.save(buf);data=buf.getvalue();media="application/vnd.openxmlformats-officedocument.presentationml.presentation"
  except ImportError:raise HTTPException(501,f"{format.upper()} export requires its optional Python dependency; install it from requirements-prod.txt")
 return Response(data,media_type=media,headers={"Content-Disposition":f'attachment; filename="{name}"',"X-Content-Type-Options":"nosniff","Cache-Control":"no-store"})
